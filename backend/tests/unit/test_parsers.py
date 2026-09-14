"""Unit tests for PDF / DOCX / text parsers."""

from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document

from app.ai.parser import parse_document_bytes
from app.ai.parser.errors import DocumentParseError
from app.modules.job.pipeline import chunk_parse_result


def _make_pdf_bytes(*pages: str) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx_bytes(*, paragraphs: list[str], heading: str | None = None) -> bytes:
    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_parse_markdown() -> None:
    raw = "# Title\n\nHello parser.\n".encode("utf-8")
    result = parse_document_bytes(raw, file_type="md")
    assert result.page_count == 1
    assert "Hello parser" in result.segments[0].text


def test_parse_pdf_keeps_pages() -> None:
    raw = _make_pdf_bytes("PageOne content AAA", "PageTwo content BBB")
    result = parse_document_bytes(raw, file_type="pdf")
    assert result.page_count == 2
    assert len(result.segments) == 2
    assert result.segments[0].page == 1
    assert "PageOne" in result.segments[0].text
    assert result.segments[1].page == 2
    assert "PageTwo" in result.segments[1].text

    chunks = chunk_parse_result(result, chunk_size=50, chunk_overlap=0)
    assert any(page == 1 for _, page, _ in chunks)
    assert any(page == 2 for _, page, _ in chunks)


def test_parse_empty_pdf_raises() -> None:
    doc = pymupdf.open()
    doc.new_page()
    raw = doc.tobytes()
    doc.close()
    with pytest.raises(DocumentParseError
) as exc:
        parse_document_bytes(raw, file_type="pdf")
    assert "扫描件" in exc.value.message or "提取文本" in exc.value.message


def test_parse_docx_with_heading_section() -> None:
    raw = _make_docx_bytes(
        heading="产品介绍",
        paragraphs=["这是正文第一段。", "这是正文第二段。"],
    )
    result = parse_document_bytes(raw, file_type="docx")
    assert result.has_text
    assert any("产品介绍" in seg.text or seg.section == "产品介绍" for seg in result.segments)
    joined = "\n".join(seg.text for seg in result.segments)
    assert "正文第一段" in joined


def test_parse_legacy_doc_bytes_rejected() -> None:
    with pytest.raises(DocumentParseError
) as exc:
        parse_document_bytes(b"not-a-zip-docx", file_type="docx")
    assert "DOCX" in exc.value.message


def test_unsupported_type() -> None:
    with pytest.raises(DocumentParseError
):
        parse_document_bytes(b"x", file_type="xlsx")

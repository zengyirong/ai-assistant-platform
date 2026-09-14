"""DOCX parser via python-docx."""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.ai.parser.base import ParseResult, ParsedSegment
from app.ai.parser.errors import DocumentParseError

_HEADING_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


def parse_docx_bytes(raw: bytes) -> ParseResult:
    if not raw:
        raise DocumentParseError("DOCUMENT_PARSE_FAILED", "DOCX 文件为空")

    # Fast reject non-zip / legacy .doc disguised as docx
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise DocumentParseError(
            "DOCUMENT_PARSE_FAILED",
            "不是有效的 DOCX（.doc 旧格式不支持，请另存为 .docx）",
        )

    try:
        document = Document(io.BytesIO(raw))
    except PackageNotFoundError as exc:
        raise DocumentParseError(
            "DOCUMENT_PARSE_FAILED",
            "无法打开 DOCX 文件",
        ) from exc
    except Exception as exc:
        raise DocumentParseError(
            "DOCUMENT_PARSE_FAILED",
            f"DOCX 解析失败：{exc}",
        ) from exc

    segments: list[ParsedSegment] = []
    buffer: list[str] = []
    current_section: str | None = None

    def flush() -> None:
        nonlocal buffer
        text = "\n".join(buffer).strip()
        buffer = []
        if text:
            segments.append(ParsedSegment(text=text, section=current_section))

    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style_name = ""
        try:
            style_name = paragraph.style.name if paragraph.style else ""
        except Exception:
            style_name = ""

        heading_match = _HEADING_RE.match(style_name or "")
        if heading_match:
            flush()
            current_section = text[:200]
            buffer.append(text)
            continue

        buffer.append(text)

    flush()

    # Tables: append as separate segments to avoid losing cell text
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [
                (cell.text or "").replace("\n", " ").strip() for cell in row.cells
            ]
            line = " | ".join(c for c in cells if c)
            if line:
                rows.append(line)
        table_text = "\n".join(rows).strip()
        if table_text:
            segments.append(
                ParsedSegment(text=table_text, section=current_section or "表格")
            )

    if not segments:
        raise DocumentParseError(
            "DOCUMENT_PARSE_FAILED",
            "DOCX 中未提取到可用正文",
        )

    return ParseResult(segments=segments, page_count=None)
